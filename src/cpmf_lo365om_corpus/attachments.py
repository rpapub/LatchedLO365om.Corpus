"""Attachment generation strategies."""

import base64
import io
import random
from mimetypes import guess_type
from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook

from .context import t
from .generators.ntfs_cases import resolve as ntfs_resolve

CONTENT_TYPE_MAP = {
    ".txt":  "text/plain",
    ".pdf":  "application/pdf",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".bin":  "application/octet-stream",
    ".png":  "image/png",
    ".jpg":  "image/jpeg",
    ".jpeg": "image/jpeg",
}


def gen_text(spec: dict, ctx: dict) -> bytes:
    content = t(spec.get("content", ""), ctx)
    return content.encode("utf-8")


def gen_pdf(spec: dict, ctx: dict) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    for page_spec in spec["pages"]:
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, t(page_spec["title"], ctx), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 12)
        pdf.ln(4)
        for line in t(page_spec["body"], ctx).split("\n"):
            pdf.cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
    buf = io.BytesIO()
    pdf.output(buf)
    return buf.getvalue()


def gen_xlsx(spec: dict, ctx: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_spec in spec["sheets"]:
        ws = wb.create_sheet(t(sheet_spec["name"], ctx))
        if "rows" in sheet_spec:
            for row in sheet_spec["rows"]:
                ws.append([t(str(cell), ctx) for cell in row])
        else:
            if "headers" in sheet_spec:
                ws.append([t(h, ctx) for h in sheet_spec["headers"]])
            template = sheet_spec.get("row_template", [])
            for i in range(sheet_spec.get("row_count", 0)):
                ws.append([t(cell, ctx, row_index=i) for cell in template])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def gen_random_bytes(spec: dict) -> bytes:
    return random.randbytes(spec["size_kb"] * 1024)


def gen_fixture(spec: dict, repo_root: Path) -> bytes:
    path = repo_root / spec["path"]
    if not path.exists():
        raise FileNotFoundError(f"fixture_file not found: {path}")
    return path.read_bytes()


def resolve_name(value, ctx: dict) -> tuple[str, str | None]:
    """Resolve an attachment name spec to (raw_name, expected_name).

    - Plain string or template → (t(value, ctx), None)
    - Strategy dict → dispatch to generator; returns (raw, expected) or (raw, None)
    """
    if isinstance(value, dict):
        strategy = value.get("strategy")
        if strategy == "ntfs_edge_case":
            raw, expected = ntfs_resolve(value)
            return raw, expected
        raise ValueError(f"Unknown name generation strategy: {strategy!r}")
    return t(value, ctx), None


def build_attachment(att_spec: dict, ctx: dict, repo_root: Path) -> dict:
    """Build a Graph API fileAttachment dict from an attachment spec."""
    strategy = att_spec["strategy"]
    name, _ = resolve_name(att_spec["name"], ctx)
    ext = Path(name).suffix.lower()
    content_type = CONTENT_TYPE_MAP.get(ext) or guess_type(name)[0] or "application/octet-stream"

    if strategy == "generate_text":
        data = gen_text(att_spec, ctx)
    elif strategy == "generate_pdf":
        data = gen_pdf(att_spec, ctx)
    elif strategy == "generate_xlsx":
        data = gen_xlsx(att_spec, ctx)
    elif strategy == "random_bytes":
        data = gen_random_bytes(att_spec)
    elif strategy == "fixture_file":
        data = gen_fixture(att_spec, repo_root)
    else:
        raise ValueError(f"Unknown attachment strategy: {strategy!r}")

    return {
        "@odata.type": "#microsoft.graph.fileAttachment",
        "name": name,
        "contentType": content_type,
        "contentBytes": base64.b64encode(data).decode("ascii"),
    }
